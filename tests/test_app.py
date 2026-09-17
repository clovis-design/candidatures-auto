import json
import smtplib
import sqlite3
import tempfile
import unittest
from concurrent.futures import ThreadPoolExecutor
from datetime import date, timedelta
from email import policy
from email.parser import Parser
from io import BytesIO
from pathlib import Path
from unittest.mock import MagicMock, patch

import openpyxl

import app as module
from tracking import TrackingStore


class ApplicationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        root = Path(self.temp.name)
        self.previous_config = module.app.config.copy()
        self.addCleanup(lambda: module.app.config.update(self.previous_config))
        module.app.config.update(TESTING=True, TRACKING_DB=str(root / 'test.sqlite3'))
        for name in ('UPLOAD_FOLDER', 'GENERATED_FOLDER'):
            patcher = patch.object(module, name, root)
            patcher.start()
            self.addCleanup(patcher.stop)
        module.jobs.clear()
        self.client = module.app.test_client()
        self.smtp = MagicMock()
        self.smtp.sendmail.return_value = {}
        patcher = patch.object(module.smtplib, 'SMTP', return_value=self.smtp)
        patcher.start()
        self.addCleanup(patcher.stop)

    def send(self, **overrides):
        data = dict(smtp_host='localhost', smtp_port='1025', smtp_user='candidate@example.org',
                    smtp_pass='', smtp_secure='none', prenom='Inès', nom='Dupont', poste='Développement',
                    emails_text='Atelier <rh@atelier.example>', personnalise='true', useCustomNames='true',
                    delay='0', subject='{ENTREPRISE} — {POSTE} — {PRENOM}',
                    lettre_template='Madame, Monsieur,\nJe candidate chez {ENTREPRISE} au poste de {POSTE}.\nL’équipe et l’innovation m’intéressent.',
                    message_body='Bonjour {ENTREPRISE}, candidature {POSTE} de {PRENOM} {NOM}.',
                    campaign_id='test-campaign')
        data.update(overrides)
        data['cv'] = (BytesIO(b'%PDF-1.4 CV de test'), 'cv.pdf')
        return self.client.post('/api/send', data=data)

    def records(self):
        return self.client.get('/api/applications').get_json()['applications']

    def workbook(self, rows, header_row=1):
        wb = openpyxl.Workbook()
        ws = wb.active
        for _ in range(header_row - 1):
            ws.append(['Titre'])
        for row in rows:
            ws.append(row)
        out = BytesIO()
        wb.save(out)
        out.seek(0)
        return out

    def test_send_persists_personalized_message_and_real_export(self):
        response = self.send()
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['success'], 1)
        message = Parser(policy=policy.default).parsestr(self.smtp.sendmail.call_args.args[2])
        self.assertEqual(str(message['Subject']), 'Atelier — Développement — Inès')
        self.assertIn('Bonjour Atelier, candidature Développement de Inès Dupont.', message.get_body().get_content())
        attachments = list(message.iter_attachments())
        self.assertEqual(len(attachments), 2)
        self.assertTrue(attachments[1].get_payload(decode=True).startswith(b'%PDF'))
        record = self.records()[0]
        self.assertEqual(record['status'], 'sent')
        self.assertEqual(record['follow_up_date'], (date.today() + timedelta(days=7)).isoformat())
        self.assertTrue(record['date_envoi'])
        wb = openpyxl.load_workbook(BytesIO(self.client.get('/api/applications/export').data))
        self.assertEqual(wb.active['A5'].value, record['date_envoi'])
        self.assertEqual(wb.active['H5'].value, 'En attente de réponse')
        self.assertFalse(list(Path(self.temp.name).glob('*.pdf')))

    def test_duplicate_and_imported_sent_dates_are_skipped(self):
        self.send()
        response = self.send(poste='développement')
        self.assertEqual(response.json['skipped'], 1)
        response = self.send(excel_entries=json.dumps([dict(email='other@example.org', date_envoi='2026-01-01')]))
        self.assertEqual(response.json['skipped'], 1)
        self.assertEqual(self.smtp.sendmail.call_count, 1)
        self.assertEqual(len(self.records()), 1)

    def test_failure_can_be_retried_without_resending_successes(self):
        self.smtp.sendmail.side_effect = [smtplib.SMTPRecipientsRefused({'rh@atelier.example': (550, b'Unknown')}), {}]
        first = self.send()
        self.assertEqual(first.json['failed'], 1)
        self.assertEqual(self.records()[0]['date_envoi'], '')
        self.assertEqual(self.records()[0]['status'], 'error')
        second = self.send()
        self.assertEqual(second.json['success'], 1)
        self.assertEqual(len(self.records()), 1)
        self.assertEqual(self.records()[0]['error'], '')

    def test_live_progress_and_duplicate_rows(self):
        counts = []
        def sendmail(*args):
            counts.append(len(self.client.get('/api/campaigns/test-campaign').json['results']))
            return {}
        self.smtp.sendmail.side_effect = sendmail
        entries = [dict(email='rh@atelier.example', stage='Python'), dict(email='rh@atelier.example', stage='Python'),
                   dict(email='rh@atelier.example', stage='Java')]
        response = self.send(excel_entries=json.dumps(entries))
        self.assertEqual(response.json['total'], 2)
        self.assertEqual(counts, [0, 1])
        self.assertTrue(self.client.get('/api/campaigns/test-campaign').json['done'])

    def test_metadata_status_notes_and_due_dates_survive_export(self):
        entries = [dict(email='rh@atelier.example', entreprise='Atelier', stage='Python',
                        interlocuteur='Camille', coordonnees='rh@atelier.example / 0102030405', notes='Premier échange')]
        self.send(excel_entries=json.dumps(entries))
        record = self.records()[0]
        self.assertEqual(record['notes'], 'Premier échange')
        response = self.client.patch(f"/api/applications/{record['id']}", json={
            'status': 'follow_up', 'follow_up_date': date.today().isoformat(),
            'date_entretien': '2026-10-20', 'notes': 'Rappeler mardi',
        })
        self.assertEqual(response.status_code, 200)
        self.assertTrue(self.records()[0]['due'])
        wb = openpyxl.load_workbook(BytesIO(self.client.get('/api/applications/export').data))
        self.assertEqual(wb.active['E5'].value, 'Camille')
        self.assertIn('0102030405', wb.active['F5'].value)
        self.assertEqual(wb.active['J5'].value, 'Rappeler mardi')
        self.assertEqual(self.client.patch(f"/api/applications/{record['id']}", json={'follow_up_date': 'wrong'}).status_code, 400)
        self.assertEqual(self.client.patch(f"/api/applications/{record['id']}", json={'status': 'error'}).status_code, 400)

    def test_flexible_excel_import_and_invalid_rows(self):
        rows = [['Entreprise', 'Email', 'Poste', 'Date entretien'],
                ['Atelier', 'rh@atelier.example', 'Python', '20/10/2026'],
                ['Atelier', 'rh@atelier.example', 'Python', ''],
                ['Atelier', 'rh@atelier.example', 'Java', ''], ['Sans email', '', '', '']]
        for header_row in (1, 4, 12):
            with self.subTest(header_row=header_row):
                response = self.client.post('/api/parse-excel', data={'file': (self.workbook(rows, header_row), 'list.xlsx')})
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response.json['count'], 2)
                self.assertEqual(len(response.json['invalid_rows']), 2)
                self.assertEqual(response.json['entries'][0]['date_entretien'], '2026-10-20')

    def test_csv_preserves_quoted_fields_and_metadata(self):
        csv = '\ufeffEntreprise;Email;Poste;Notes\n"Atelier; Ouest";rh@atelier.example;Python;"Bonjour, Camille"\n'
        response = self.client.post('/api/parse-excel', data={'file': (BytesIO(csv.encode()), 'list.csv')})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json['entries'][0]['entreprise'], 'Atelier; Ouest')
        self.assertEqual(response.json['entries'][0]['notes'], 'Bonjour, Camille')

    def test_draft_export_has_no_fake_send_date_and_no_formulas(self):
        response = self.client.post('/api/generate-excel', json={'entries': [dict(email='rh@atelier.example', entreprise='=1+1', stage='Python')], 'poste': 'Java'})
        wb = openpyxl.load_workbook(BytesIO(response.data))
        self.assertIsNone(wb.active['A5'].value)
        self.assertEqual(wb.active['H5'].value, 'À envoyer')
        self.assertEqual(wb.active['B5'].data_type, 's')
        self.assertEqual(wb.active['D5'].value, 'Python')

    def test_update_original_workbook_keeps_other_sheets_and_formulas(self):
        self.send()
        stream = self.workbook([['Email', 'Poste', 'Calcul'], ['rh@atelier.example', 'Développement', '=1+2'], ['unknown@example.org', 'Python', 42]])
        wb = openpyxl.load_workbook(stream)
        wb.active['C2'].font = openpyxl.styles.Font(bold=True)
        wb.create_sheet('À conserver')['A1'] = 'Autres données'
        upload = BytesIO(); wb.save(upload); upload.seek(0)
        response = self.client.post('/api/update-excel', data={'file': (upload, 'original.xlsx')})
        self.assertEqual(response.headers['X-Updated-Rows'], '1')
        result = openpyxl.load_workbook(BytesIO(response.data))
        self.assertEqual(result.active['C2'].value, '=1+2')
        self.assertTrue(result.active['C2'].font.bold)
        self.assertEqual(result['À conserver']['A1'].value, 'Autres données')
        self.assertEqual(result.active['E2'].value, 'En attente de réponse')
        self.assertIsNone(result.active['D3'].value)

    def test_atomic_reservation_across_connections(self):
        path = module.app.config['TRACKING_DB']
        TrackingStore(path)
        def reserve(_):
            return TrackingStore(path).reserve(dict(email='same@example.org', stage='Python'))
        with ThreadPoolExecutor(max_workers=4) as executor:
            results = list(executor.map(reserve, range(4)))
        self.assertEqual(sum(r is not None for r in results), 1)

    def test_smtp_test_does_not_send_and_auth_errors_are_not_ignored(self):
        self.smtp.login.side_effect = smtplib.SMTPAuthenticationError(535, b'AUTH failed')
        response = self.client.post('/api/test-smtp', json=dict(smtp_host='localhost', smtp_port='1025',
            smtp_user='candidate@example.org', smtp_pass='secret', smtp_secure='none'))
        self.assertEqual(response.status_code, 400)
        self.smtp.sendmail.assert_not_called()
        self.smtp.close.assert_called()
        self.assertEqual(self.records(), [])

    def test_invalid_delay_and_invalid_spreadsheet_are_actionable(self):
        for delay in ('NaN', '-1', 'infinity', 'abc', '61'):
            self.assertEqual(self.send(delay=delay).status_code, 400)
        self.smtp.sendmail.assert_not_called()
        response = self.client.post('/api/parse-excel', data={'file': (self.workbook([['Inconnu'], ['Texte']]), 'bad.xlsx')})
        self.assertEqual(response.status_code, 400)
        self.assertIn('En-têtes introuvables', response.json['error'])
        self.assertEqual(self.client.get('/').status_code, 200)

    def test_auto_letter_with_browser_boolean_and_unicode_profile(self):
        response = self.client.post('/api/generate-letter', json={
            'prenom': 'Inès', 'nom': 'Dupont', 'poste': 'Développement',
            'personnalise': True, 'experience': 'L’équipe développe des œuvres numériques…',
        })
        self.assertEqual(response.status_code, 200)
        self.assertIn('{ENTREPRISE_NOM}', response.json['template'])
        self.assertIn('{POSTE}', response.json['template'])
        self.assertNotIn('{POSTE}', response.json['preview'])
        pdf = self.client.get(response.json['pdf_url'])
        self.assertTrue(pdf.data.startswith(b'%PDF'))
        pdf.close()

    def test_accepted_email_is_not_retried_if_tracking_save_fails(self):
        with patch.object(TrackingStore, 'finish', side_effect=OSError('Disque indisponible')):
            response = self.send()
        self.assertEqual(response.json['failed'], 1)
        self.assertIn('Email accepté', response.json['results'][0]['error'])
        self.assertEqual(self.records()[0]['status'], 'sending')
        self.assertEqual(self.send().json['skipped'], 1)
        self.assertEqual(self.smtp.sendmail.call_count, 1)

    def test_manual_applications_without_email_persist_and_can_be_edited(self):
        data = dict(entreprise='Atelier', stage='Python', status='sent', date_envoi='2026-09-15',
                    source='LinkedIn', url='https://example.org/jobs/42', notes='Déposée sur le site')
        first = self.client.post('/api/applications', json=data)
        self.assertEqual(first.status_code, 201)
        self.assertEqual(self.client.post('/api/applications', json={**data, 'entreprise': 'Studio'}).status_code, 201)
        records = TrackingStore(module.app.config['TRACKING_DB']).all()
        self.assertEqual(len(records), 2)
        self.assertTrue(all(r['email'] == '' for r in records))
        self.assertEqual(self.client.get('/api/applications').json['stats']['sent'], 2)
        changed = self.client.patch(f"/api/applications/{first.json['id']}", json={
            'status': 'interview', 'date_entretien': '2026-09-22', 'source': 'Site carrière',
            'entreprise': 'Atelier Numérique', 'notes': 'Entretien prévu',
        })
        self.assertEqual(changed.status_code, 200)
        record = next(r for r in self.records() if r['id'] == first.json['id'])
        self.assertEqual(record['source'], 'Site carrière')
        self.assertEqual(record['url'], data['url'])
        self.assertEqual(record['date_envoi'], '15/09/2026')
        self.assertTrue(record['editable'])
        self.smtp.sendmail.assert_not_called()

    def test_manual_validation_and_duplicates_do_not_overwrite_history(self):
        base = dict(entreprise='Atelier', stage='Python', status='sent')
        invalid = [dict(entreprise=''), dict(stage=''), dict(email='invalide'),
                   dict(url='javascript:alert(1)'), dict(date_envoi='31/02/2026'),
                   dict(status='inconnu'), dict(status='draft', date_envoi='2026-01-01')]
        for change in invalid:
            with self.subTest(change=change):
                self.assertEqual(self.client.post('/api/applications', json={**base, **change}).status_code, 400)
        self.assertEqual(self.records(), [])
        response = self.client.post('/api/applications', json={**base, 'notes': 'À conserver'})
        self.assertEqual(response.status_code, 201)
        duplicate = self.client.post('/api/applications', json={**base, 'entreprise': 'ATELIER', 'stage': 'python', 'notes': 'Écraser'})
        self.assertEqual(duplicate.status_code, 409)
        self.assertEqual(self.records()[0]['notes'], 'À conserver')
        # Le poste doit garder sa ponctuation : C++ et C# sont distincts.
        for stage in ('C++', 'C#'):
            self.assertEqual(self.client.post('/api/applications', json={**base, 'stage': stage}).status_code, 201)
        another = self.client.post('/api/applications', json={**base, 'entreprise': 'Studio'})
        conflict = self.client.patch(f"/api/applications/{another.json['id']}", json={'entreprise': 'Atelier'})
        self.assertEqual(conflict.status_code, 409)
        self.smtp.sendmail.assert_not_called()

    def test_import_history_mixed_rows_reports_errors_and_ignores_duplicates(self):
        rows = [['Entreprise', 'Poste', 'Email', 'Date candidature', 'Statut', 'Site', 'Lien de l’offre', 'Notes'],
                ['Atelier', 'Python', '', '15/09/2026', 'Entretien', 'LinkedIn', 'https://example.org/42', 'Contact Camille'],
                ['Studio', 'Python', 'rh@studio.example', '', '', 'Indeed', '', 'À préparer'],
                ['Atelier', 'Python', '', '15/09/2026', 'Entretien', '', '', 'Doublon'],
                ['Erreur date', 'Java', '', '31/02/2026', 'En attente', '', '', ''],
                ['Erreur statut', 'Java', '', '', 'Mystère', '', '', ''],
                ['Date inconnue', 'Web', '', '', 'En attente', '', '', '']]
        def upload():
            return self.client.post('/api/applications/import', data={'file': (self.workbook(rows, 4), 'suivi.xlsx')})
        first = upload()
        self.assertEqual(first.status_code, 200)
        self.assertEqual((first.json['created'], first.json['skipped'], first.json['invalid']), (3, 1, 2))
        self.assertEqual([r['row'] for r in first.json['invalid_rows']], [8, 9])
        records = {r['entreprise']: r for r in self.records()}
        self.assertEqual(records['Atelier']['source'], 'LinkedIn')
        self.assertEqual(records['Atelier']['url'], 'https://example.org/42')
        self.assertEqual(records['Atelier']['notes'], 'Contact Camille')
        self.assertEqual(records['Studio']['status'], 'draft')
        self.assertEqual(records['Studio']['sent_at'], '')
        self.assertEqual(records['Date inconnue']['status'], 'sent')
        self.assertEqual(records['Date inconnue']['sent_at'], '')
        self.assertEqual(self.client.get('/api/applications').json['stats']['sent'], 2)
        self.client.patch(f"/api/applications/{records['Atelier']['id']}", json={'notes': 'Modifiée dans le suivi'})
        second = upload()
        self.assertEqual((second.json['created'], second.json['skipped'], second.json['invalid']), (0, 4, 2))
        self.assertIn('Modifiée dans le suivi', [r['notes'] for r in self.records()])
        self.smtp.sendmail.assert_not_called()

    def test_history_csv_without_email_and_campaign_import_remains_strict(self):
        content = 'Entreprise;Poste;Statut;Source;Date relance\n"Atelier; Ouest";Python;À relancer;Indeed;01/09/2026\n'
        response = self.client.post('/api/applications/import', data={'file': (BytesIO(content.encode()), 'suivi.csv')})
        self.assertEqual(response.json['created'], 1)
        self.assertEqual(self.records()[0]['entreprise'], 'Atelier; Ouest')
        self.assertTrue(self.records()[0]['due'])
        campaign = self.client.post('/api/parse-excel', data={'file': (BytesIO(content.encode()), 'suivi.csv')})
        self.assertEqual(campaign.status_code, 400)
        invalid = self.client.post('/api/applications/import', data={'file': (BytesIO(b'invalid'), 'suivi.xlsx')})
        self.assertEqual(invalid.status_code, 400)

    def test_external_sent_status_blocks_resending_even_without_known_date(self):
        self.client.post('/api/applications', json=dict(entreprise='Atelier', stage='Développement', status='sent'))
        response = self.send()
        self.assertEqual(response.json['skipped'], 1)
        self.smtp.sendmail.assert_not_called()
        self.assertEqual(len(self.records()), 1)

    def test_imported_draft_can_be_updated_then_sent_without_duplicate_record(self):
        response = self.client.post('/api/applications', json=dict(entreprise='Atelier', stage='Python', status='draft'))
        record_id = response.json['id']
        changed = self.client.patch(f'/api/applications/{record_id}', json={'stage': 'Développement', 'notes': 'À conserver'})
        self.assertEqual(changed.status_code, 200)
        self.assertEqual(self.send().json['success'], 1)
        self.assertEqual(len(self.records()), 1)
        self.assertEqual(self.records()[0]['id'], record_id)
        self.assertEqual(self.records()[0]['origin'], 'smtp')
        self.assertEqual(self.records()[0]['email'], 'rh@atelier.example')
        self.assertEqual(self.records()[0]['notes'], 'À conserver')

    def test_existing_database_migration_keeps_ids_and_results(self):
        path = module.app.config['TRACKING_DB']
        with sqlite3.connect(path) as db:
            db.execute('''CREATE TABLE applications (id INTEGER PRIMARY KEY, email TEXT NOT NULL,
                stage TEXT NOT NULL, details TEXT NOT NULL, status TEXT NOT NULL, sent_at TEXT DEFAULT '',
                updated_at TEXT NOT NULL, follow_up_date TEXT DEFAULT '', notes TEXT DEFAULT '', error TEXT DEFAULT '',
                UNIQUE(email, stage))''')
            db.execute('INSERT INTO applications VALUES (?,?,?,?,?,?,?,?,?,?)',
                       (42, 'old@example.org', 'Python', json.dumps({'entreprise': 'Ancienne'}), 'accepted',
                        '2026-09-01T10:00:00', '2026-09-01T10:00:00', '', 'Conserver', ''))
        response = self.client.post('/api/applications', json=dict(entreprise='Nouvelle', stage='Python', status='sent'))
        self.assertEqual(response.status_code, 201)
        old = next(r for r in self.records() if r['id'] == 42)
        self.assertEqual(old['notes'], 'Conserver')
        self.assertEqual(old['status'], 'accepted')
        self.assertEqual(old['origin'], 'smtp')
        self.assertEqual(len(TrackingStore(path).all()), 2)

    def test_external_export_round_trip_preserves_source_url_email_and_notes(self):
        self.client.post('/api/applications', json=dict(entreprise='Atelier', stage='Python', email='rh@atelier.example',
            coordonnees='0102030405', source='LinkedIn', url='https://example.org/42', notes='=1+1', status='interview',
            date_envoi='2026-09-15', date_entretien='2026-10-01'))
        export = self.client.get('/api/applications/export')
        wb = openpyxl.load_workbook(BytesIO(export.data))
        self.assertEqual(wb.active['K5'].value, 'LinkedIn')
        self.assertEqual(wb.active['L5'].value, 'https://example.org/42')
        self.assertEqual(wb.active['J5'].data_type, 's')
        self.assertIn('rh@atelier.example', wb.active['F5'].value)
        self.assertIn('0102030405', wb.active['F5'].value)
        module.app.config['TRACKING_DB'] = str(Path(self.temp.name) / 'roundtrip.sqlite3')
        response = self.client.post('/api/applications/import', data={'file': (BytesIO(export.data), 'suivi.xlsx')})
        self.assertEqual(response.json['created'], 1)
        record = self.records()[0]
        self.assertEqual(record['source'], 'LinkedIn')
        self.assertEqual(record['date_entretien'], '2026-10-01')
        self.assertEqual(record['notes'], '=1+1')

    def test_update_workbook_matches_external_company_and_position_without_email(self):
        self.client.post('/api/applications', json=dict(entreprise='Atelier', stage='Python', status='interview'))
        upload = self.workbook([['Entreprise', 'Poste', 'Statut'], ['Atelier', 'Python', 'En attente'], ['Studio', 'Python', 'En attente']])
        response = self.client.post('/api/update-excel', data={'file': (upload, 'suivi.xlsx')})
        self.assertEqual(response.headers['X-Updated-Rows'], '1')
        wb = openpyxl.load_workbook(BytesIO(response.data))
        self.assertEqual(wb.active['C2'].value, 'Entretien')
        self.assertEqual(wb.active['C3'].value, 'En attente')


if __name__ == '__main__':
    unittest.main()
