"""Import de tableurs usuels et mise à jour d'un classeur existant."""
import csv
import re
from datetime import date, datetime
from io import BytesIO, StringIO

import openpyxl

EMAIL = re.compile(r'[a-zA-Z0-9_.+%\-]+@[a-zA-Z0-9-]+(?:\.[a-zA-Z0-9-]+)+')
KEYS = ['date_envoi', 'entreprise', 'type_candidature', 'stage', 'interlocuteur',
        'coordonnees', 'date_entretien', 'resultats', 'follow_up_date', 'notes']


def cell_text(value):
    if isinstance(value, (date, datetime)):
        return value.strftime('%Y-%m-%d')
    return str(value).strip() if value is not None else ''


def detect_headers(rows, header_keys, normalize):
    aliases = {
        'email': 'coordonnees', 'mail': 'coordonnees', 'adresseemail': 'coordonnees',
        'emailcontact': 'coordonnees', 'societe': 'entreprise', 'company': 'entreprise',
        'poste': 'stage', 'postevise': 'stage', 'contact': 'interlocuteur',
        'statut': 'resultats', 'status': 'resultats', 'type': 'type_candidature',
        'relance': 'follow_up_date', 'daterelance': 'follow_up_date', 'notes': 'notes',
    }
    best = None
    for number, row in enumerate(rows[:20], 1):
        mapping = {}
        for col, value in enumerate(row, 1):
            norm = normalize(value)
            key = header_keys.get(norm) or aliases.get(norm)
            if key and key not in mapping.values():
                mapping[col] = key
        if 'coordonnees' in mapping.values() and (best is None or len(mapping) > len(best[1])):
            best = number, mapping
    if not best:
        raise ValueError('En-têtes introuvables : ajoutez une colonne Email ou Coordonnées dans les 20 premières lignes.')
    return best


def parse_rows(rows, header_keys, normalize, company_name):
    header_row, mapping = detect_headers(rows, header_keys, normalize)
    entries, invalid, seen = [], [], set()
    for number, row in enumerate(rows[header_row:], header_row + 1):
        if not any(cell_text(value) for value in row):
            continue
        entry = {key: cell_text(row[col - 1]) if col <= len(row) else '' for col, key in mapping.items()}
        for key in ('date_envoi', 'date_entretien', 'follow_up_date'):
            value = entry.get(key, '')
            if value.lower() in ('aucun', 'aucune', '-', 'non'):
                entry[key] = ''
            elif value:
                for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%Y-%m-%d %H:%M:%S'):
                    try:
                        entry[key] = datetime.strptime(value, fmt).date().isoformat()
                        break
                    except ValueError:
                        continue
        matches = EMAIL.findall(entry.get('coordonnees', ''))
        if len(matches) != 1:
            invalid.append({'row': number, 'reason': 'Un seul email est requis dans la colonne Email / Coordonnées.'})
            continue
        entry['email'] = matches[0].lower()
        entry.setdefault('stage', '')
        identity = (entry['email'], entry['stage'].casefold())
        if identity in seen:
            invalid.append({'row': number, 'reason': 'Doublon email / poste ignoré.'})
            continue
        seen.add(identity)
        entry['entreprise'] = entry.get('entreprise') or company_name(entry['email'])
        entry['source_row'] = number
        entries.append(entry)
    headers = [(col, cell_text(rows[header_row - 1][col - 1]), key) for col, key in mapping.items()]
    return entries, invalid, headers


def read_excel(stream, header_keys, normalize, company_name):
    wb = openpyxl.load_workbook(stream, data_only=True, read_only=True)
    try:
        ws = wb.active
        if ws.max_row and ws.max_row > 10000 or ws.max_column and ws.max_column > 200:
            raise ValueError('Le tableur est limité à 10 000 lignes et 200 colonnes.')
        rows = list(ws.iter_rows(values_only=True))
        return parse_rows(rows, header_keys, normalize, company_name)
    finally:
        wb.close()


def read_csv(stream, header_keys, normalize, company_name):
    content = stream.read().decode('utf-8-sig')
    try:
        dialect = csv.Sniffer().sniff(content[:8192], delimiters=';,\t')
    except csv.Error:
        dialect = None
    rows = list(csv.reader(StringIO(content), dialect)) if dialect else list(csv.reader(StringIO(content), delimiter=';'))
    if len(rows) > 10000:
        raise ValueError('Le tableur est limité à 10 000 lignes.')
    return parse_rows(rows, header_keys, normalize, company_name)


def write_text(cell, value):
    """Les données importées restent du texte, jamais des formules exécutables."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = 's'


def update_workbook(stream, records, header_keys, normalize):
    wb = openpyxl.load_workbook(stream)
    ws = wb.active
    if ws.max_row > 10000 or ws.max_column > 200:
        raise ValueError('Le tableur est limité à 10 000 lignes et 200 colonnes.')
    rows = list(ws.iter_rows(values_only=True))
    header_row, mapping = detect_headers(rows, header_keys, normalize)
    columns = {key: col for col, key in mapping.items()}
    for key, title in [('date_envoi', 'Date envoi'), ('resultats', 'Résultats'),
                       ('date_entretien', 'Date entretien'), ('follow_up_date', 'Date relance'), ('notes', 'Notes')]:
        if key not in columns:
            col = ws.max_column + 1
            columns[key] = col
            ws.cell(header_row, col, title)
    count = 0
    for number in range(header_row + 1, ws.max_row + 1):
        matches = EMAIL.findall(cell_text(ws.cell(number, columns['coordonnees']).value))
        if len(matches) != 1:
            continue
        stage = cell_text(ws.cell(number, columns['stage']).value) if 'stage' in columns else ''
        candidates = [r for r in records if r['email'] == matches[0].lower() and
                      (not stage or r['stage'].casefold() == stage.casefold())]
        if len(candidates) != 1:
            continue  # Ne pas deviner lorsque plusieurs postes correspondent.
        record = candidates[0]
        for key in ('date_envoi', 'resultats', 'date_entretien', 'follow_up_date', 'notes'):
            write_text(ws.cell(number, columns[key]), record.get(key, ''))
        count += 1
    output = BytesIO()
    wb.save(output)
    wb.close()
    output.seek(0)
    return output, count
